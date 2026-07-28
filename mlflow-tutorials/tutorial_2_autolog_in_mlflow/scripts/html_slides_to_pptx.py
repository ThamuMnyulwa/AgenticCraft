from __future__ import annotations

import base64
import re
import sys
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BLUE = RGBColor(0x42, 0x85, 0xF4)
RED = RGBColor(0xEA, 0x43, 0x35)
YELLOW = RGBColor(0xFB, 0xBC, 0x04)
GREEN = RGBColor(0x34, 0xA8, 0x53)
INK = RGBColor(0x20, 0x21, 0x24)
GREY = RGBColor(0x5F, 0x63, 0x68)
LINE = RGBColor(0xDA, 0xDC, 0xE0)
CODE_BG = RGBColor(0xF1, 0xF3, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENTS = [BLUE, RED, YELLOW, GREEN]

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_X = Inches(1.05)
TOP_Y = Inches(0.65)


def clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\xa0", " ")).strip()


def add_text(slide, text, x, y, w, h, *, size=22, color=INK, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    return box


def slide_number(slide, n: int) -> None:
    add_text(slide, f"{n:02d}", Inches(11.8), Inches(0.35), Inches(0.6), Inches(0.25), size=10, color=GREY, bold=True, align=PP_ALIGN.RIGHT)


def brand_dots(slide, y=Inches(0.95)) -> None:
    for i, color in enumerate(ACCENTS):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN_X + i * Inches(0.26), y, Inches(0.19), Inches(0.19))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()


def eyebrow(slide, text: str, accent=BLUE) -> None:
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, MARGIN_X, TOP_Y, Inches(0.12), Inches(0.12))
    dot.fill.solid()
    dot.fill.fore_color.rgb = accent
    dot.line.fill.background()
    add_text(slide, text.upper(), MARGIN_X + Inches(0.22), TOP_Y - Inches(0.03), Inches(8.5), Inches(0.28), size=10, color=GREY, bold=True)


def title(slide, text: str, y=Inches(1.12), size=34) -> None:
    add_text(slide, text, MARGIN_X, y, Inches(10.7), Inches(0.8), size=size, bold=True)


def lead(slide, text: str, y=Inches(1.95), h=Inches(0.75)) -> None:
    if text:
        add_text(slide, text, MARGIN_X, y, Inches(8.0), h, size=17, color=GREY)


def code_block(slide, code: str, x, y, w, h, *, accent=BLUE, size=11) -> None:
    bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    bg.fill.solid()
    bg.fill.fore_color.rgb = CODE_BG
    bg.line.fill.background()
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.07), h)
    marker.fill.solid()
    marker.fill.fore_color.rgb = accent
    marker.line.fill.background()
    box = slide.shapes.add_textbox(x + Inches(0.22), y + Inches(0.16), w - Inches(0.42), h - Inches(0.25))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(code.rstrip().splitlines()):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        r.font.name = "Courier New"
        r.font.size = Pt(size)
        r.font.color.rgb = INK


def card(slide, x, y, w, h, title_text: str, body: str, *, tag=None, accent=BLUE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = LINE
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.055), h)
    marker.fill.solid()
    marker.fill.fore_color.rgb = accent
    marker.line.fill.background()
    ty = y + Inches(0.25)
    if tag:
        add_text(slide, tag.upper(), x + Inches(0.28), ty, w - Inches(0.5), Inches(0.2), size=8, color=accent, bold=True)
        ty += Inches(0.28)
    add_text(slide, title_text, x + Inches(0.28), ty, w - Inches(0.5), Inches(0.3), size=15, bold=True)
    add_text(slide, body, x + Inches(0.28), ty + Inches(0.42), w - Inches(0.5), h - Inches(0.75), size=11, color=GREY)


def cards_grid(slide, cards, y=Inches(2.45)) -> None:
    gap_x, gap_y = Inches(0.35), Inches(0.32)
    w, h = (Inches(11.2) - gap_x) / 2, Inches(1.35)
    for i, item in enumerate(cards):
        x = MARGIN_X + (i % 2) * (w + gap_x)
        cy = y + (i // 2) * (h + gap_y)
        card(slide, x, cy, w, h, item["title"], item["body"], tag=item.get("tag"), accent=ACCENTS[i % 4])


def bullets(slide, items, y=Inches(2.25)) -> None:
    for i, item in enumerate(items):
        cy = y + i * Inches(0.72)
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN_X, cy + Inches(0.04), Inches(0.07), Inches(0.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENTS[i % 4]
        bar.line.fill.background()
        box = slide.shapes.add_textbox(MARGIN_X + Inches(0.22), cy, Inches(10.4), Inches(0.65))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        if item["bold"]:
            r = p.add_run()
            r.text = item["bold"] + " "
            r.font.bold = True
            r.font.name = "Aptos"
            r.font.size = Pt(16)
            r.font.color.rgb = INK
        r = p.add_run()
        r.text = item["rest"]
        r.font.name = "Aptos"
        r.font.size = Pt(16)
        r.font.color.rgb = INK


def callout(slide, text: str, y=Inches(5.95), w=Inches(8.8)) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN_X, y, w, Inches(0.72))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.fill.background()
    add_text(slide, text, MARGIN_X + Inches(0.28), y + Inches(0.12), w - Inches(0.55), Inches(0.5), size=12)


def logo_path(section) -> Path | None:
    img = section.find("img")
    if not img or not img.get("src", "").startswith("data:image/png;base64,"):
        return None
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    tmp.write(base64.b64decode(img["src"].split(",", 1)[1]))
    tmp.close()
    return Path(tmp.name)


def section_data(section):
    heading = section.find(["h1", "h2"])
    lead_el = section.select_one("p.lead")
    eyebrow_el = section.select_one(".eyebrow")
    card_data = []
    for c in section.select(".card"):
        card_data.append(
            {
                "title": clean(c.find("h3").get_text(" ", strip=True)) if c.find("h3") else "",
                "body": clean(c.find("p").get_text(" ", strip=True)) if c.find("p") else "",
                "tag": clean(c.select_one(".tag").get_text(" ", strip=True)) if c.select_one(".tag") else None,
            }
        )
    bullet_data = []
    for li in section.select("ul.points li"):
        b = clean(li.find("b").get_text(" ", strip=True)) if li.find("b") else ""
        full = clean(li.get_text(" ", strip=True))
        bullet_data.append({"bold": b, "rest": clean(full[len(b) :]) if b and full.startswith(b) else full})
    return {
        "title": clean(heading.get_text(" ", strip=True)) if heading else "",
        "lead": clean(lead_el.get_text(" ", strip=True)) if lead_el else "",
        "eyebrow": clean(eyebrow_el.get_text(" ", strip=True)) if eyebrow_el else "",
        "cards": card_data,
        "bullets": bullet_data,
        "pre": section.find_all("pre"),
        "callout": clean(section.select_one(".callout").get_text(" ", strip=True)) if section.select_one(".callout") else "",
    }


def convert(html_path: Path, pptx_path: Path) -> None:
    soup = BeautifulSoup(html_path.read_text(), "html.parser")
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for n, section in enumerate(soup.select("section.slide"), 1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        slide_number(slide, n)
        data = section_data(section)

        if n == 1:
            brand_dots(slide)
            if path := logo_path(section):
                slide.shapes.add_picture(str(path), MARGIN_X, Inches(1.45), width=Inches(3.05))
            add_text(slide, data["title"], MARGIN_X, Inches(3.6), Inches(8.8), Inches(0.55), size=26, color=GREY)
            lead(slide, data["lead"], y=Inches(4.6))
            continue

        if n == 13:
            brand_dots(slide)
            add_text(slide, data["title"].replace(". ", ".\n"), MARGIN_X, Inches(1.65), Inches(9.0), Inches(1.5), size=34, bold=True)
            lead(slide, data["lead"], y=Inches(3.5))
            if data["pre"]:
                code_block(slide, data["pre"][0].get_text(), MARGIN_X, Inches(4.65), Inches(6.0), Inches(0.75), accent=YELLOW, size=14)
            continue

        eyebrow(slide, data["eyebrow"] or f"Slide {n}", ACCENTS[(n - 1) % 4])
        title(slide, data["title"])
        lead(slide, data["lead"])

        if data["cards"]:
            cards_grid(slide, data["cards"], y=Inches(3.05 if n == 12 else 2.55))
            if data["callout"]:
                callout(slide, data["callout"], y=Inches(6.15), w=Inches(9.9))
        elif data["bullets"]:
            bullets(slide, data["bullets"], y=Inches(2.35 if data["lead"] else 2.15))
        elif n == 7 and len(data["pre"]) >= 2:
            add_text(slide, "Before", MARGIN_X, Inches(2.05), Inches(4.2), Inches(0.28), size=12, color=RED, bold=True)
            add_text(slide, "After", Inches(7.05), Inches(2.05), Inches(4.2), Inches(0.28), size=12, color=GREEN, bold=True)
            code_block(slide, data["pre"][0].get_text(), MARGIN_X, Inches(2.42), Inches(5.55), Inches(3.75), accent=RED, size=10)
            code_block(slide, data["pre"][1].get_text(), Inches(7.05), Inches(2.42), Inches(4.7), Inches(2.0), accent=GREEN, size=12)
        elif data["pre"]:
            h = Inches(3.85 if n == 8 else 3.35 if n == 9 else 2.55)
            size = 9 if n in {8, 9} else 12
            code_block(slide, data["pre"][0].get_text(), MARGIN_X, Inches(2.65), Inches(8.5), h, size=size)
            if data["callout"]:
                callout(slide, data["callout"])

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx_path)


def main() -> int:
    html_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("slides/mlflow-slides.html")
    pptx_path = Path(sys.argv[2]) if len(sys.argv) > 2 else html_path.with_suffix(".pptx")
    convert(html_path, pptx_path)
    print(pptx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
